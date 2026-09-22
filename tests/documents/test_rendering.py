"""PDF / DOCX rendering: validity, content equivalence, layout sanity, determinism."""

import io

import pytest
from docx import Document
from pypdf import PdfReader

from app.documents.content import cover_letter_model, normalize_text, resume_model
from app.documents.models import (
    Contact,
    ContactLink,
    DocumentFormat,
    DocumentModel,
    DocumentValidation,
    RenderInput,
    Section,
)
from app.documents.renderers.base import ascii_safe
from app.documents.renderers.docx import DocxRenderer
from app.documents.renderers.pdf import PdfRenderer
from app.documents.storage import sha256_bytes
from app.documents.validate import extract_text, validate_document
from tests.documents.conftest import ready_preparation

CONTACT = Contact(name="Ribhu Siripurapu", email="ribhu@example.com", phone="+91 90000 00000", location="Hyderabad, India", links=[ContactLink(label="GitHub", url="https://github.com/ribhu")])


def _input(model, fmt):
    return RenderInput(tenant_id="t", preparation_id="p", preparation_version=1, preparation_fingerprint="f", artifact_type=model.artifact_type, format=fmt, model=model)


def _resume(bullets=3):
    return DocumentModel(
        artifact_type="RESUME",
        contact=CONTACT,
        sections=[
            Section(key="summary", title="Summary", items=[("Backend engineer — Python/Go, distributed systems… ships ticketing engines.", False)]),
            Section(key="skills", title="Skills", items=[("Skills: Python, Go, FastAPI, PostgreSQL", True)]),
            Section(key="projects", title="Projects", items=[(f"Ticket Engine {i}: booking system handling 12,000 req/s at p99 45 ms. Technologies: Go, Redis.", True) for i in range(bullets)]),
            Section(key="education", title="Education", items=[("B.Tech CSE (Data Science), MGIT, graduating 2027, CGPA 8.9", True)]),
        ],
    )


@pytest.mark.parametrize("fmt, renderer", [(DocumentFormat.PDF, PdfRenderer()), (DocumentFormat.DOCX, DocxRenderer())])
def test_render_is_valid_extractable_ordered_and_complete(fmt, renderer):
    model = _resume()
    result = renderer.render(_input(model, fmt))
    assert result.data[:4] in (b"%PDF", b"PK\x03\x04")
    pages, texts = extract_text(result.data, fmt)
    text = normalize_text("\n".join(texts))
    assert pages == 1
    # Every line of content appears, and the sections keep their order.
    for line in model.expected_lines():  # headings are upper-cased by layout, nothing else changes
        assert normalize_text(line).replace(" ", "").lower() in text.replace(" ", "").lower(), line
    positions = [text.find(title.upper()) for title in ("Summary", "Skills", "Projects", "Education")]
    assert positions == sorted(positions) and positions[0] >= 0
    assert "12,000" in text and "2027" in text and "8.9" in text and "https://github.com/ribhu" in text
    outcome = validate_document(result.data, fmt, model)
    assert outcome.status is DocumentValidation.PASSED and outcome.missing_lines == []


def test_cover_letter_renders_paragraphs_subject_and_signature():
    model = DocumentModel(artifact_type="COVER_LETTER", contact=CONTACT, subject="Re: Backend Engineer at Acme", paragraphs=["Dear Acme Hiring Team,", "I am applying for the Backend Engineer position at Acme.", "Thank you for your consideration."], signature="Ribhu Siripurapu")
    for fmt, renderer in ((DocumentFormat.PDF, PdfRenderer()), (DocumentFormat.DOCX, DocxRenderer())):
        result = renderer.render(_input(model, fmt))
        assert validate_document(result.data, fmt, model).status is DocumentValidation.PASSED
        _, texts = extract_text(result.data, fmt)
        assert "Re: Backend Engineer at Acme" in normalize_text(texts[0]) and "Thank you for your consideration." in normalize_text(texts[0])


def test_pdf_links_are_clickable_and_metadata_fixed():
    result = PdfRenderer().render(_input(_resume(), DocumentFormat.PDF))
    reader = PdfReader(io.BytesIO(result.data))
    annots = reader.pages[0].get("/Annots") or []
    uris = [a.get_object()["/A"]["/URI"] for a in annots if "/A" in a.get_object()]
    assert "https://github.com/ribhu" in uris
    assert reader.metadata.get("/Producer", "").startswith("CareerOS")
    assert "2000" in str(reader.metadata.get("/CreationDate", ""))


def test_docx_opens_and_keeps_hyperlink_and_bullets():
    result = DocxRenderer().render(_input(_resume(), DocumentFormat.DOCX))
    document = Document(io.BytesIO(result.data))
    bullets = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 5  # skills + 3 projects + education
    rels = [r.target_ref for r in document.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://github.com/ribhu" in rels


@pytest.mark.parametrize("fmt, renderer", [(DocumentFormat.PDF, PdfRenderer()), (DocumentFormat.DOCX, DocxRenderer())])
def test_rendering_is_byte_deterministic(fmt, renderer):
    a = renderer.render(_input(_resume(), fmt)).data
    b = renderer.render(_input(_resume(), fmt)).data
    assert sha256_bytes(a) == sha256_bytes(b)
    changed = _resume()
    changed.sections[1].items[0] = ("Skills: Python, Go, Rust", True)
    assert sha256_bytes(renderer.render(_input(changed, fmt)).data) != sha256_bytes(a)


def test_multi_page_resume_gets_page_numbers_and_oversized_is_flagged():
    big = _resume(bullets=120)
    result = PdfRenderer().render(_input(big, DocumentFormat.PDF))
    assert result.page_count >= 2
    _, texts = extract_text(result.data, DocumentFormat.PDF)
    assert f"Page 1 of {result.page_count}" in texts[0]
    outcome = validate_document(result.data, DocumentFormat.PDF, big, max_pages=1)
    assert outcome.status is DocumentValidation.NEEDS_REVIEW and any("exceeds" in i for i in outcome.issues)
    assert outcome.missing_lines == [], "flagged for layout, but no content was lost"


def test_validation_detects_lost_content_and_malformed_files():
    model = _resume()
    outcome = validate_document(b"%PDF-1.4 garbage", DocumentFormat.PDF, model)
    assert outcome.status is DocumentValidation.FAILED and "malformed" in outcome.issues[0]
    other = _resume()
    other.sections[3].items[0] = ("M.Sc. Physics, 2031", True)
    rendered = PdfRenderer().render(_input(other, DocumentFormat.PDF)).data
    outcome = validate_document(rendered, DocumentFormat.PDF, model)
    assert outcome.status is DocumentValidation.FAILED and outcome.missing_lines == ["B.Tech CSE (Data Science), MGIT, graduating 2027, CGPA 8.9"]
    assert validate_document(b"", DocumentFormat.PDF, model).status is DocumentValidation.FAILED


def test_empty_model_is_refused():
    from app.documents.renderers.base import RenderError

    empty = DocumentModel(artifact_type="RESUME", contact=Contact())
    for renderer in (PdfRenderer(), DocxRenderer()):
        assert not renderer.can_render(_input(empty, renderer.format))
        with pytest.raises(RenderError):
            renderer.render(_input(empty, renderer.format))


def test_ascii_fallback_only_touches_typography():
    assert ascii_safe("Go — 12,000 req/s… “fast”") == 'Go - 12,000 req/s... "fast"'


def test_unicode_content_embeds_a_font_and_keeps_every_glyph():
    from app.documents.renderers.pdf import needs_unicode_font

    latin = _resume()
    assert not needs_unicode_font(latin), "typographic punctuation alone does not force embedding"
    extended = _resume()
    extended.sections[3].items[0] = ("B.Tech CSE, MGIT — advisor Łukasz Ćwikła, graduating 2027", True)
    assert needs_unicode_font(extended)
    renderer = PdfRenderer()
    if renderer.font_path is None:
        pytest.skip("no TrueType font available on this machine")
    assert renderer.font_for(latin) is None and renderer.font_for(extended) == renderer.font_path
    result = renderer.render(_input(extended, DocumentFormat.PDF))
    outcome = validate_document(result.data, DocumentFormat.PDF, extended)
    assert outcome.status is DocumentValidation.PASSED, outcome.missing_lines
    # A script the font does not cover is caught, never silently rendered as boxes.
    devanagari = _resume()
    devanagari.sections[3].items[0] = ("B.Tech CSE, MGIT (महात्मा गांधी), graduating 2027", True)
    result = renderer.render(_input(devanagari, DocumentFormat.PDF))
    outcome = validate_document(result.data, DocumentFormat.PDF, devanagari)
    assert outcome.status is DocumentValidation.FAILED and outcome.missing_lines


def test_models_from_real_preparation_blocks(db_session, tenant_id, opportunities, documents):
    prep = ready_preparation(db_session, tenant_id, opportunities)
    resume = documents.model_for(prep, "RESUME")
    letter = documents.model_for(prep, "COVER_LETTER")
    source = {a.artifact_type: a for a in prep.artifacts}
    # Every block's text is in the model, in order; nothing added beyond identity + job facts.
    block_texts = [" ".join(b["text"].split()) for b in source["RESUME"].blocks]
    model_texts = [t for s in resume.sections for t, _ in s.items]
    assert model_texts == block_texts
    assert resume.contact.email == "ribhu@example.com" and resume.contact.links
    assert letter is not None and letter.paragraphs[0].startswith("Dear") and letter.signature
    assert letter.subject and "Backend Engineer" in letter.subject
    assert cover_letter_model([], resume.contact, "X", "Y") is None
    assert resume_model([{"section": "skills", "text": "   "}], resume.contact).sections == []
