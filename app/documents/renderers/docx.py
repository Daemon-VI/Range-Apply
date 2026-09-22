"""DOCX renderer (python-docx) from the same DocumentModel as the PDF.

Determinism: python-docx stamps zip entries with the current time and core
properties with now(); both are normalised (fixed timestamps) so identical
input yields identical bytes.
"""

import io
import zipfile
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.documents.models import DocumentFormat, DocumentModel, RenderInput, RenderResult
from app.documents.renderers.base import Renderer, RenderError

_FIXED_DATE = datetime(2000, 1, 1, tzinfo=timezone.utc)
_ZIP_DATE = (1980, 1, 1, 0, 0, 0)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "143C8C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    node.set(qn("xml:space"), "preserve")
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


class DocxRenderer(Renderer):
    name = "python-docx"
    version = "docx-v1"
    format = DocumentFormat.DOCX

    def can_render(self, render_input: RenderInput) -> bool:
        return render_input.format is DocumentFormat.DOCX and bool(render_input.model.expected_lines())

    def render(self, render_input: RenderInput) -> RenderResult:
        model = render_input.model
        if not model.expected_lines():
            raise RenderError("nothing to render: the preparation has no content")
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        for section in document.sections:
            section.top_margin = section.bottom_margin = Pt(50)
            section.left_margin = section.right_margin = Pt(54)
        self._header(document, model)
        if model.artifact_type == "RESUME":
            self._resume(document, model)
        else:
            self._letter(document, model)
        props = document.core_properties
        props.author = model.contact.name or "CareerOS"
        props.title = f"{model.contact.name} - {'Resume' if model.artifact_type == 'RESUME' else 'Cover letter'}".strip(" -")
        props.created = _FIXED_DATE.replace(tzinfo=None)
        props.modified = _FIXED_DATE.replace(tzinfo=None)
        props.last_modified_by = "CareerOS"
        props.revision = 1
        buffer = io.BytesIO()
        document.save(buffer)
        data = _normalise_zip(buffer.getvalue())
        return RenderResult(data=data, page_count=1)

    def _header(self, document, model: DocumentModel) -> None:
        contact = model.contact
        if contact.name:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(contact.name)
            run.bold = True
            run.font.size = Pt(17)
        bits = [b for b in (contact.email, contact.phone, contact.location) if b]
        if bits:
            paragraph = document.add_paragraph("  |  ".join(bits))
            paragraph.runs[0].font.size = Pt(9.5)
        if contact.links:
            paragraph = document.add_paragraph()
            for index, link in enumerate(contact.links):
                _add_hyperlink(paragraph, link.url, link.url)
                if index < len(contact.links) - 1:
                    paragraph.add_run("  |  ")
        rule = document.add_paragraph()
        rule.paragraph_format.space_after = Pt(2)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "A0A0A0")):
            bottom.set(qn(key), value)
        border.append(bottom)
        rule._p.get_or_add_pPr().append(border)

    def _resume(self, document, model: DocumentModel) -> None:
        for section in model.sections:
            heading = document.add_paragraph()
            run = heading.add_run(section.title.upper())
            run.bold = True
            run.font.size = Pt(11)
            heading.paragraph_format.space_before = Pt(8)
            for text, bulleted in section.items:
                if bulleted:
                    document.add_paragraph(text, style="List Bullet")
                else:
                    document.add_paragraph(text)

    def _letter(self, document, model: DocumentModel) -> None:
        if model.date_line:
            document.add_paragraph(model.date_line)
        if model.subject:
            paragraph = document.add_paragraph()
            paragraph.add_run(model.subject).bold = True
        for text in model.paragraphs:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if model.signature:
            paragraph = document.add_paragraph(model.signature)
            paragraph.runs[0].font.color.rgb = RGBColor(0, 0, 0)


def _normalise_zip(data: bytes) -> bytes:
    """Rewrite the package with fixed entry timestamps so bytes are reproducible."""
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in sorted(source.infolist(), key=lambda i: i.filename):
            fixed = zipfile.ZipInfo(info.filename, date_time=_ZIP_DATE)
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            target.writestr(fixed, source.read(info.filename))
    return out.getvalue()
