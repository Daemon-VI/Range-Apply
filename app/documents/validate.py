"""Structural and content validation of rendered documents.

Structure: the file parses, has pages, no empty page, page count within
the configured bound, size within bound. Content: every expected line of
the DocumentModel is found in the extracted text after normalisation, so a
renderer can never silently drop or alter candidate content.
"""

import io
from typing import Optional

from app.config import settings
from app.documents.content import normalize_text
from app.documents.models import (
    DocumentFormat,
    DocumentModel,
    DocumentValidation,
    ValidationOutcome,
)


def extract_text(data: bytes, fmt: DocumentFormat) -> tuple[int, list[str]]:
    """(page count, text per page). DOCX has no pages; its text is one 'page'."""
    if fmt is DocumentFormat.PDF:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        return len(pages), pages
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.append(cell.text)
    # Hyperlink runs live outside paragraph.text; collect their text too.
    from docx.oxml.ns import qn

    for hyperlink in document.element.body.iter(qn("w:hyperlink")):
        lines.extend(node.text or "" for node in hyperlink.iter(qn("w:t")))
    return 1, ["\n".join(lines)]


def validate_document(data: bytes, fmt: DocumentFormat, model: DocumentModel, max_pages: Optional[int] = None) -> ValidationOutcome:
    issues: list[str] = []
    if not data:
        return ValidationOutcome(status=DocumentValidation.FAILED, issues=["empty output"])
    if len(data) > settings.documents_max_bytes:
        issues.append(f"output is {len(data)} bytes, above the {settings.documents_max_bytes}-byte bound")
    try:
        page_count, pages = extract_text(data, fmt)
    except Exception as exc:  # noqa: BLE001 - a parse failure is the finding
        return ValidationOutcome(status=DocumentValidation.FAILED, issues=[f"malformed {fmt.value}: {type(exc).__name__}: {str(exc)[:120]}"])
    if page_count == 0:
        return ValidationOutcome(status=DocumentValidation.FAILED, issues=["zero pages"])
    empty_pages = [i + 1 for i, text in enumerate(pages) if not text.strip()]
    if empty_pages:
        issues.append(f"empty page(s): {empty_pages}")
    limit = max_pages or (settings.documents_max_pages_resume if model.artifact_type == "RESUME" else settings.documents_max_pages_cover_letter)
    if page_count > limit:
        issues.append(f"{page_count} pages exceeds the {limit}-page bound for {model.artifact_type}")
    extracted = normalize_text("\n".join(pages))
    haystack = _squash(extracted)
    missing = [line for line in model.expected_lines() if _squash(normalize_text(line)) not in haystack]
    if model.artifact_type == "RESUME" and not model.sections:
        issues.append("resume has no sections")
    if model.artifact_type == "COVER_LETTER" and not model.paragraphs:
        issues.append("cover letter has no paragraphs")
    if missing:
        return ValidationOutcome(status=DocumentValidation.FAILED, page_count=page_count, issues=issues + [f"{len(missing)} expected line(s) missing from the rendered text"], missing_lines=missing[:10], extracted_chars=len(extracted))
    status = DocumentValidation.NEEDS_REVIEW if issues else DocumentValidation.PASSED
    return ValidationOutcome(status=status, page_count=page_count, issues=issues, extracted_chars=len(extracted))


def _squash(text: str) -> str:
    """Comparison form that ignores line-wrap artefacts (spaces, hyphenation)."""
    return "".join(ch for ch in text.lower() if ch.isalnum())
